#!/usr/bin/env python3
"""
gen_testplan.py — turn an IP's machine-readable vplan.yaml + requirements.md into a
traditional, human-facing verification test plan as an Excel (.xlsx) workbook.

This bridges the framework's YAML vPlan (great for git/automation/agents) to the
industry-traditional spreadsheet testplan (feature -> test case -> checker ->
coverage point -> requirement), the form DV leads actually review and sign off.

Usage:  python3 flow/scripts/gen_testplan.py ips/<ip>
Output: ips/<ip>/vplan/<ip>_testplan.xlsx   (Test Plan / Traceability / Coverage sheets)

Deps: openpyxl, pyyaml.
"""
import re, sys
from pathlib import Path
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F6F5C")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
SEC_FILL = PatternFill("solid", fgColor="E4EEE9")
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILL = {"covered": "E4EEE9", "closed": "CDECDD", "waived": "F5E4DC", "planned": "F5F0D0"}


def parse_requirements(md: Path):
    """Pull REQ rows out of the requirements markdown table."""
    reqs = {}
    if not md.exists():
        return reqs
    for line in md.read_text(encoding="utf-8", errors="replace").splitlines():
        m = re.match(r"\|\s*(REQ-[A-Z]+-\d+)\s*\|(.+)", line)
        if not m:
            continue
        rid = m.group(1)
        cells = [c.strip() for c in m.group(2).split("|")]
        reqs[rid] = {"text": cells[0] if cells else "",
                     "category": cells[1] if len(cells) > 1 else "",
                     "priority": cells[2] if len(cells) > 2 else ""}
    return reqs


def checker_for(item):
    """Derive the pass/fail checker mechanism from method + coverage kind."""
    kind = (item.get("coverage") or {}).get("kind", "")
    method = item.get("method", "")
    if kind == "assertion":
        return f"SVA assertion: {item['coverage'].get('name','')}"
    if method == "register" or kind == "ral":
        return "RAL predictor + uvm_reg sequences (reset / bit-bash / access)"
    if method == "static":
        return "Static gate (Verible lint / Yosys elaboration)"
    return "Reference-model scoreboard + SVA"


def coverage_points_for(item):
    cov = item.get("coverage") or {}
    kind = cov.get("kind", "")
    if kind == "covergroup":
        cps = cov.get("coverpoints") or {}
        pts = ", ".join(cps.keys()) if isinstance(cps, dict) else str(cps)
        cross = cov.get("cross")
        s = f"{cov.get('name','')} [{pts}]"
        if cross:
            s += f" x({'/'.join(cross) if isinstance(cross,list) else cross})"
        return s
    if kind == "assertion":
        return f"assertion cover: {cov.get('name','')}"
    if kind == "ral":
        return "RAL: " + ", ".join(cov.get("ral_sequences", []))
    if kind == "alias":
        return "errata roll-up (aliases existing bins): " + ", ".join((cov.get("aliases") or {}).keys())
    if kind == "static":
        return "code-coverage / elaboration gate"
    return kind or "-"


def regression_result(reports: Path):
    f = reports / "regression.txt"
    if not f.exists():
        return "not run"
    head = f.read_text(errors="replace").splitlines()[:1]
    m = re.search(r"(\d+)/(\d+)\s+runs?\s+passed", head[0]) if head else None
    return f"{m.group(1)}/{m.group(2)} runs passed" if m else "see regression.txt"


def xcrg_functional_score(reports: Path):
    html = reports / "_cov" / "functional_report" / "functionalCoverageReport" / "dashboard.html"
    if not html.exists():
        return None
    t = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.read_text(errors="replace")))
    m = re.search(r"Coverage Summary Score Inst Score ([0-9.]+)", t)
    return f"{float(m.group(1)):.1f}%" if m else None


def xcrg_code_scores(reports: Path):
    html = reports / "_cov" / "code_report" / "codeCoverageReport" / "dashboard.html"
    if not html.exists():
        return {}
    t = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", html.read_text(errors="replace")))
    # dashboard lists the 4 score LABELS, then counts (files/modules/instances) + the
    # 4 SCORES as a run of 7 numbers: files modules instances stmt branch cond toggle
    m = re.search(r"Statement Coverage Score Branch Coverage Score Condition Coverage Score "
                  r"Toggle Coverage Score\s+" + r"([0-9.]+)\s+" * 6 + r"([0-9.]+)", t)
    if not m:
        return {}
    stmt, branch, cond, tog = m.group(4), m.group(5), m.group(6), m.group(7)
    return {"statement": f"{float(stmt):.1f}%", "branch": f"{float(branch):.1f}%",
            "condition": f"{float(cond):.1f}%", "toggle": f"{float(tog):.1f}%"}


def code_waivers(ip_dir: Path):
    """List the module/instance/signal scoping waivers from the IP's exclusion file."""
    ip = ip_dir.name
    f = ip_dir / "dv" / f"{ip}_cov_exclusions.txt"
    if not f.exists():
        return []
    out = []
    for line in f.read_text(errors="replace").splitlines():
        line = line.strip()
        if line and not line.startswith("#"):
            out.append(line)
    return out


def set_cols(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HDR_FILL, HDR_FONT, WRAP, BORDER
    ws.freeze_panes = "A2"


def style_rows(ws, status_col=None):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
            c.border = BORDER
        if status_col:
            sc = row[status_col - 1]
            fill = STATUS_FILL.get(str(sc.value).lower())
            if fill:
                sc.fill = PatternFill("solid", fgColor=fill)


def main():
    if len(sys.argv) < 2:
        sys.exit("usage: gen_testplan.py ips/<ip>")
    ip_dir = Path(sys.argv[1]).resolve()
    ip = ip_dir.name
    vplan = yaml.safe_load((ip_dir / "vplan" / "vplan.yaml").read_text(encoding="utf-8"))
    items = vplan.get("items", [])
    reqs = parse_requirements(ip_dir / "spec" / "requirements.md")

    wb = Workbook()

    # ---- Sheet 1: Test Plan ----
    ws = wb.active
    ws.title = "Test Plan"
    header(ws, ["VP-ID", "Requirement(s)", "Feature / Description", "Test Type",
                "Test Case(s)", "Checker (pass/fail)", "Coverage Point(s)", "Status"])
    for it in items:
        ws.append([
            it.get("id", ""),
            ", ".join(it.get("traces_to", [])),
            it.get("feature", ""),
            it.get("method", ""),
            ", ".join(it.get("tests", [])),
            checker_for(it),
            coverage_points_for(it),
            it.get("status", "planned"),
        ])
    set_cols(ws, [16, 20, 52, 18, 34, 34, 40, 10])
    style_rows(ws, status_col=8)

    # ---- Sheet 2: Requirements Traceability ----
    ws2 = wb.create_sheet("Requirements Traceability")
    header(ws2, ["Requirement", "Description", "Priority", "vPlan Item(s)", "Test(s)", "Status"])
    for rid, r in reqs.items():
        covering = [it for it in items if rid in it.get("traces_to", [])]
        vps = ", ".join(it["id"] for it in covering)
        tsts = sorted({t for it in covering for t in it.get("tests", [])})
        if not covering:
            status = "NO vPlan ITEM"
        else:
            sset = {it.get("status", "planned") for it in covering}
            status = "covered" if sset <= {"covered", "closed", "waived"} else "open"
        ws2.append([rid, r["text"], r.get("priority", ""), vps or "—", ", ".join(tsts) or "—", status])
    set_cols(ws2, [16, 70, 8, 30, 40, 14])
    style_rows(ws2, status_col=6)

    # ---- Sheet 3: Coverage Summary (DATA-DRIVEN — no hardcoded numbers) ----
    ws3 = wb.create_sheet("Coverage Summary")
    ws3.append([f"{ip.upper()} — Coverage & Waivers"]); ws3["A1"].font = Font(bold=True, size=13)
    ws3.append([f"goal: {vplan.get('coverage_goal_pct','?')}% functional coverage"])
    ws3.append([])
    reports = ip_dir / "reports"
    rows = [["Metric", "Result", "Source"]]
    rows.append(["Seeded regression", regression_result(reports), "reports/regression.txt"])
    fscore = xcrg_functional_score(reports)
    rows.append(["Functional coverage (union)", fscore or "not run", "reports/_cov/functional_report/"])
    code = xcrg_code_scores(reports)
    if code:
        for k in ("statement", "branch", "condition", "toggle"):
            if k in code:
                rows.append([f"DUT code — {k}", code[k], "reports/_cov/code_report/ (DUT-scoped)"])
    else:
        rows.append(["DUT code coverage", "not run", "reports/_cov/code_report/"])
    rows.append([])
    rows.append(["Code-coverage waivers (module scoping)", "Type", "Source"])
    for m in code_waivers(ip_dir):
        rows.append([m, "code scoping", f"dv/{ip}_cov_exclusions.txt"])
    if not code_waivers(ip_dir):
        rows.append(["(no exclusion file present)", "", ""])
    for r in rows:
        ws3.append(r)
        if r and r[0] in ("Metric", "Code-coverage waivers (module scoping)"):
            for c in ws3[ws3.max_row]:
                c.fill, c.font = SEC_FILL, Font(bold=True)
    set_cols(ws3, [42, 34, 46])
    for row in ws3.iter_rows(min_row=4):
        for c in row:
            c.alignment = WRAP

    out = ip_dir / "vplan" / f"{ip}_testplan.xlsx"
    wb.save(out)
    print(f"wrote {out}  ({len(items)} test-plan items, {len(reqs)} requirements)")


if __name__ == "__main__":
    main()
